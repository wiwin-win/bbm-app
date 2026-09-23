"""Selftest aplikasi Monitoring BBM & Alat (jalan tanpa server).

    python selftest.py [file_uji.xlsx]

Memakai DB sementara di folder test sehingga data asli tidak tersentuh.
"""
import io
import os
import re
import sys
import tempfile
from datetime import date
from pathlib import Path

BASE = Path(__file__).resolve().parent
TMPL = BASE / "templates"
TMP = Path(tempfile.mkdtemp(prefix="bbm-test-"))
os.environ["BBM_DB"] = str(TMP / "uji.db")
sys.path.insert(0, str(BASE))

import openpyxl  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

import app as aplikasi  # noqa: E402
import excel  # noqa: E402
import report  # noqa: E402
import storage as db  # noqa: E402

LULUS = GAGAL = 0


def cek(nama: str, syarat: bool, info: str = "") -> None:
    global LULUS, GAGAL
    if syarat:
        LULUS += 1
        print(f"  [OK]   {nama}")
    else:
        GAGAL += 1
        print(f"  [GAGAL] {nama} {info}")


def main() -> int:
    berkas = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(r"C:\Users\wiwin\Documents\BBM.xlsx")
    print(f"DB uji : {os.environ['BBM_DB']}")
    db.init_db()

    # ---------- 1. master + entry dasar
    print("\n1) Master & entry")
    db.import_master("unit", [{"no_lambung": "LOADER (GAS 009)"}, {"no_lambung": "TANGKI BBM"}])
    db.import_master("driver", [{"nama": "RONALD"}])
    db.import_master("location", [{"nama": "PT GAS"}])
    db.import_master("activity", [{"nama": "MAINTENANCE JALAN"}])
    cek("master unit tersimpan", any(u["no_lambung"] == "LOADER (GAS 009)"
                                     for u in db.master_list("unit")))
    cek("master driver tersimpan", any(u["nama"] == "RONALD" for u in db.master_list("driver")))
    db.set_settings({"stok_awal": "1000", "harga_default": "17200"})

    id1 = db.entry_simpan({"tanggal": "2026-09-10", "driver": "RONALD",
                           "no_lambung": "LOADER (GAS 009)", "aktivitas": "MAINTENANCE JALAN",
                           "lokasi": "PT GAS", "keluar": 150, "harga": 17200}, user_id=1)
    id2 = db.entry_simpan({"tanggal": "2026-09-11", "driver": "RONALD",
                           "no_lambung": "TANGKI BBM", "aktivitas": "BBM MASUK",
                           "lokasi": "PT GAS", "masuk": 4900, "harga": 17200}, user_id=1)
    id3 = db.entry_simpan({"tanggal": "2026-09-12", "driver": "RONALD",
                           "no_lambung": "LOADER (GAS 009)", "lokasi": "PT GAS",
                           "keluar": 100, "harga": 17200}, user_id=1)
    d = db.entry_daftar({}, limit=50)
    cek("3 entry tersimpan", d["total"] == 3, f"dapat {d['total']}")
    cek("total keluar = 250", abs(d["total_keluar"] - 250) < 0.001, str(d["total_keluar"]))
    cek("total masuk = 4900", abs(d["total_masuk"] - 4900) < 0.001, str(d["total_masuk"]))

    # ---------- 2. sisa stok berjalan
    print("\n2) Sisa stok berjalan")
    rows = db.semua_entries({})
    cek("sisa stok akhir = 1000+4900-250 = 5650", abs(rows[-1]["sisa_stok"] - 5650) < 0.001,
        str(rows[-1]["sisa_stok"]))
    cek("sisa stok harian urut", abs(rows[0]["sisa_stok"] - 850) < 0.001, str(rows[0]["sisa_stok"]))
    # sisa stok harus benar walau limit kecil (dulu dihitung setelah limit -> ngawur)
    kecil = db.entry_daftar({}, limit=1)
    cek("sisa stok baris terbaru benar walau limit=1",
        abs(kecil["rows"][0]["sisa_stok"] - rows[-1]["sisa_stok"]) < 0.001,
        f'limit={kecil["rows"][0]["sisa_stok"]} penuh={rows[-1]["sisa_stok"]}')
    dua = db.entry_daftar({}, limit=2)["rows"]
    cek("sisa stok konsisten antar-2 baris teratas",
        abs(dua[0]["sisa_stok"] - (dua[1]["sisa_stok"] + dua[0]["masuk"] - dua[0]["keluar"])) < 0.001,
        f'{dua[0]["sisa_stok"]} vs {dua[1]["sisa_stok"]}')
    db.set_settings({"stok_awal": "5000"})
    cek("sisa stok ikut naik saat stok awal diubah",
        abs(db.entry_daftar({}, limit=1)["rows"][0]["sisa_stok"] - 9650) < 0.001,
        str(db.entry_daftar({}, limit=1)["rows"][0]["sisa_stok"]))
    db.set_settings({"stok_awal": "1000"})

    # ---------- 3. filter & rekap
    print("\n3) Filter & rekap")
    cek("filter lokasi", db.entry_daftar({"lokasi": "PT GAS"})["total"] == 3)
    cek("filter alat", db.entry_daftar({"alat": "LOADER"})["total"] == 2)
    cek("filter bulan", db.entry_daftar({"bulan": "2026-09"})["total"] == 3)
    cek("filter bulan lain kosong", db.entry_daftar({"bulan": "2026-08"})["total"] == 0)
    rk = db.rekap({})
    cek("rekap biaya = 250*17200", abs(rk["total_biaya"] - 250 * 17200) < 1, str(rk["total_biaya"]))
    cek("rekap per_alat ada 2", len(rk["per_alat"]) == 2)
    cek("rekap per_lokasi ada 1", len(rk["per_lokasi"]) == 1)

    # ---------- 4. ubah & hapus
    print("\n4) Ubah & hapus")
    db.entry_simpan({"tanggal": "2026-09-10", "no_lambung": "LOADER (GAS 009)",
                     "lokasi": "PT GAS", "keluar": 200, "harga": 17200}, rec_id=id1)
    cek("data terubah", db.entry_ambil(id1)["keluar"] == 200)
    db.entry_hapus(id3)
    cek("data terhapus", db.entry_ambil(id3) is None)
    cek("jumlah jadi 2", db.entry_daftar({})["total"] == 2)

    # ---------- 5. cetak laporan
    print("\n5) Laporan")
    f = {"bulan": "2026-09"}
    teks = report.laporan_teks(db.rekap(f), db.get_settings(), f)
    cek("laporan teks memuat judul", "LAPORAN BBM" in teks)
    cek("laporan teks memuat sisa stok", "Sisa stok" in teks)
    html = report.laporan_html(db.rekap(f), db.semua_entries(f), db.get_settings(), f)
    cek("laporan html memuat tabel", "<table" in html and "Cetak / Simpan PDF" in html)

    # ---------- 6. export excel
    print("\n6) Export Excel")
    bio = excel.ekspor_excel(db.semua_entries(f), db.rekap(f), db.get_settings())
    wb = openpyxl.load_workbook(io.BytesIO(bio.getvalue()))
    cek("sheet BBM ada", "BBM" in wb.sheetnames)
    cek("sheet REKAP ada", "REKAP" in wb.sheetnames)
    ws = wb["BBM"]
    cek("judul di A1", "PENCATATAN" in str(ws["A1"].value))
    cek("header TANGGAL ada", any(str(c.value).upper() == "TANGGAL"
                                  for c in ws[3] if c.value))
    # posisi kolom WAJIB sama antara header dan baris data (dulu pernah bergeser
    # karena header dan isi ditulis dari dua daftar terpisah)
    hdr = {str(c.value).strip().upper(): c.column for c in ws[3] if c.value}
    cek("kolom KELUAR & MASUK tidak tertukar",
        hdr.get("BBM / LTR (KELUAR)", 0) < hdr.get("BBM MASUK (LTR)", 999))
    cek("kolom MASUK & HARGA tidak tertukar",
        hdr.get("BBM MASUK (LTR)", 0) < hdr.get("HARGA / LTR (RP)", 999))
    cek("kolom KETERANGAN di posisi terakhir (sama seperti Excel asli)",
        hdr.get("KETERANGAN", 0) == max(hdr.values()))
    baris_data = [c.value for c in ws[4]]
    cek("nilai masuk ada di kolom BBM MASUK",
        baris_data[hdr["BBM MASUK (LTR)"] - 1] in (0, None)
        or isinstance(baris_data[hdr["BBM MASUK (LTR)"] - 1], (int, float)))
    # angka di kolom keluar harus sama dengan isi DB (bukan angka kolom sebelah)
    r0 = db.semua_entries(f)[0]
    cek("angka kolom KELUAR sama dengan DB",
        abs((baris_data[hdr["BBM / LTR (KELUAR)"] - 1] or 0) - r0["keluar"]) < 0.01)
    cek("angka kolom HARGA sama dengan DB",
        abs((baris_data[hdr["HARGA / LTR (RP)"] - 1] or 0) - r0["harga"]) < 0.01)

    # ---------- 7. import excel asli
    print("\n7) Import Excel")
    cek("_tgl_ketat tolak teks", excel._tgl_ketat("BBM HILANG/BOCOR") is None)
    cek("_tgl_ketat terima datetime", excel._tgl_ketat(__import__("datetime").datetime(2026, 5, 12))
        == "2026-05-12")
    cek("_tgl_ketat terima serial", excel._tgl_ketat(46154) == "2026-05-12")
    cek("_tgl_ketat terima teks ISO", excel._tgl_ketat("2026-05-12") == "2026-05-12")
    if berkas.exists():
        hasil = excel.impor_excel(str(berkas))
        cek("import membaca baris > 100", hasil["baris"] > 100, str(hasil["baris"]))
        cek("import menandai 1 catatan khusus",
            len(hasil["catatan_khusus"]) == 1, str(hasil["catatan_khusus"]))
        cek("catatan khusus BBM HILANG 3290 L",
            hasil["catatan_khusus"] and abs(hasil["catatan_khusus"][0]["keluar"] - 3290) < 0.01)
        cek("import mengisi master driver", len(db.master_list("driver")) > 100)
        cek("import mengisi master alat", len(db.master_list("unit")) > 100)
        tgl_ok = all(len(r["tanggal"]) == 10 and r["tanggal"][4] == "-"
                     for r in db.semua_entries({})[:50])
        cek("tanggal ternormalisasi YYYY-MM-DD", tgl_ok)
        cek("semua tanggal valid", not [r for r in db.semua_entries({}) if len(r["tanggal"]) != 10])
        rk2 = db.rekap({})
        cek("rekap hasil import punya biaya", rk2["total_biaya"] > 0, str(rk2["total_biaya"]))
        cek("tidak ada baris TOTAL / SISA STOK AKHIR",
            not any(x["no_lambung"].upper() in ("TOTAL", "SISA STOK AKHIR", "SISA STOK AKHIR (LTR)")
                    for x in db.semua_entries({})))
        cek("baris 'BBM HILANG/BOCOR' tetap tersimpan sebagai catatan",
            any("BBM HILANG" in x["no_lambung"].upper() and abs(x["keluar"] - 3290) < 0.01
                for x in db.semua_entries({})))
        cek("keterangan catatan khusus terisi",
            any("catatan khusus dari Excel" in (x["keterangan"] or "")
                for x in db.semua_entries({})))
        cek("export excel tidak error walau ada tanggal kosong", True)
        _b = excel.ekspor_excel([{"tanggal": "", "driver": "", "no_lambung": "X", "aktivitas": "",
                                  "lokasi": "", "keluar": 1, "masuk": 0, "harga": 1, "keterangan": ""}],
                                rk2, db.get_settings())
        cek("export excel tahan tanggal tidak valid", _b.getvalue()[:2] == b"PK")
    else:
        print("  (file uji tidak ada, dilewati)")

    # ---------- 8. template excel
    print("\n8) Template Excel")
    wt = openpyxl.load_workbook(io.BytesIO(excel.template_kosong().getvalue()))
    cek("template punya 11 kolom", wt.active.max_column == 11,
        str(wt.active.max_column))

    # ---------- 9. HTTP end-to-end
    print("\n9) API HTTP")
    with TestClient(aplikasi.app) as cl:
        cek("halaman / tanpa login redirect", cl.get("/", follow_redirects=False).status_code == 307)
        salah = cl.post("/api/login", data={"username": "admin", "password": "salah"})
        cek("login salah ditolak", salah.status_code == 401)
        masuk = cl.post("/api/login", data={"username": "admin", "password": "admin"})
        cek("login admin berhasil", masuk.status_code == 200 and masuk.json()["ok"])
        cek("halaman / terbuka", cl.get("/").status_code == 200)
        cek("halaman /m terbuka", cl.get("/m").status_code == 200)
        r = cl.get("/api/saya").json()
        cek("role admin", r["user"]["role"] == "admin")
        r = cl.get("/api/master").json()
        cek("api master mengembalikan unit", len(r["unit"]) > 0)
        r = cl.get("/api/entries?limit=5").json()
        cek("api entries limit 5", len(r["rows"]) == 5, str(len(r["rows"])))
        r = cl.get("/api/ringkasan").json()
        cek("api ringkasan ada stok", "sisa_stok" in r["data"])
        r = cl.get("/api/rekap").json()
        cek("api rekap ada per_alat", len(r["data"]["per_alat"]) > 0)
        baru = cl.post("/api/entries", json={"tanggal": "2026-09-20", "no_lambung": "UJI API",
                                            "lokasi": "PT GAS", "keluar": 10})
        cek("tambah via API", baru.status_code == 200)
        rid = baru.json()["id"]
        kosong = cl.post("/api/entries", json={"tanggal": "2026-09-20", "no_lambung": "X"})
        cek("tolak entry tanpa liter", kosong.status_code == 400)
        cek("ubah via API", cl.put(f"/api/entries/{rid}", json={
            "tanggal": "2026-09-20", "no_lambung": "UJI API", "keluar": 12}).status_code == 200)
        cek("hapus via API", cl.delete(f"/api/entries/{rid}").status_code == 200)
        cek("hapus id tidak ada tetap OK", cl.delete("/api/entries/999999").status_code == 200)
        x = cl.get("/api/export/excel?bulan=2026-09")
        cek("export xlsx via API", x.status_code == 200 and len(x.content) > 5000)
        cek("content-type xlsx", "spreadsheetml" in x.headers["content-type"])
        cek("template xlsx via API", cl.get("/api/template/excel").status_code == 200)
        wa = cl.get("/api/laporan/wa?bulan=2026-09").json()
        cek("laporan WA via API", "LAPORAN BBM" in wa["teks"])
        cek("laporan html via API", cl.get("/api/laporan/html").status_code == 200)
        cek("users ada admin", any(u["username"] == "admin" for u in cl.get("/api/users").json()["rows"]))
        cek("setting tersimpan", cl.post("/api/settings", json={"nama_instansi": "UJI PT"}).status_code == 200)
        cek("setting terbaca", cl.get("/api/saya").json()["settings"]["nama_instansi"] == "UJI PT")
        cek("backup database", cl.get("/api/backup").json()["ok"])
        cek("log aktivitas terisi", len(cl.get("/api/log").json()["rows"]) > 3)
        cek("manifest PWA", cl.get("/manifest.webmanifest").status_code == 200)
        cek("sw.js tersaji", cl.get("/sw.js").status_code == 200)
        cek("ikon PWA ada", cl.get("/static/icon-192.png").status_code == 200)
        # operator: tidak boleh hapus
        cl.post("/api/users", json={"username": "op1", "password": "rahasia", "role": "operator"})
        cl2 = TestClient(aplikasi.app)
        cl2.post("/api/login", data={"username": "op1", "password": "rahasia"})
        cek("operator diblokir hapus", cl2.delete("/api/entries/1").status_code == 403)
        cek("operator diblokir admin", cl2.get("/api/users").status_code == 403)
        cek("operator boleh input", cl2.post("/api/entries", json={
            "tanggal": "2026-09-21", "no_lambung": "OP TEST", "keluar": 5}).status_code == 200)
        cek("logout", cl2.post("/api/logout").status_code == 200)
        # anti data dobel: kirim 2x dengan kunci sama -> hanya 1 baris
        k1 = cl.post("/api/entries", json={"tanggal": "2026-09-22", "no_lambung": "UJI DOBEL",
                                           "keluar": 7, "kunci": "hp-uji-123"})
        k2 = cl.post("/api/entries", json={"tanggal": "2026-09-22", "no_lambung": "UJI DOBEL",
                                           "keluar": 7, "kunci": "hp-uji-123"})
        cek("kunci sama tidak digandakan", k1.json()["id"] == k2.json()["id"],
            f'{k1.json()["id"]} vs {k2.json()["id"]}')
        beda = cl.post("/api/entries", json={"tanggal": "2026-09-22", "no_lambung": "UJI DOBEL",
                                            "keluar": 7, "kunci": "hp-uji-456"})
        cek("kunci beda tetap masuk", beda.json()["id"] != k1.json()["id"])
        cek("kunci tersimpan di baris", any(r.get("kunci") == "hp-uji-123"
            for r in cl.get("/api/entries?alat=UJI DOBEL").json()["rows"]))

    # ---------- 10. halaman HP: input + tarik Excel
    print("\n10) Halaman HP (input + tarik Excel)")
    import pathlib
    tmpl = {n: pathlib.Path(TMPL / n).read_text(encoding="utf-8")
            for n in ("mobile.html", "index.html", "login.html")}
    mob = tmpl["mobile.html"]
    cek("tab Excel ada", 'data-t="xl"' in mob and 'id="vXl"' in mob)
    cek("tombol Tarik Data", "tarikExcel(false)" in mob)
    cek("tombol Bagikan/Kirim", "tarikExcel(true)" in mob)
    cek("pakai navigator.share untuk kirim file", "navigator.share" in mob
        and "canShare" in mob and "files:[file]" in mob.replace(" ", ""))
    cek("ada fallback unduh file", "a.download" in mob or "a.download=nama" in mob.replace(" ", ""))
    cek("buat Excel di HP (xlsx.js ter-inline/dimuat)", "BBMXlsx" in mob or "xlsx.js" in mob)
    cek("anti data dobel pakai kunci unik", "kunciBaru()" in mob and "kunci:_kunci" in mob)
    cek("nama file Excel dibuat di HP", "namaBerkas" in mob)
    cek("kirim ringkasan WhatsApp", "kirimWA" in mob and "navigator.share" in mob)
    cek("data selalu disimpan di HP dulu", "a.push(o); setDataLokal(a)" in mob.replace("  ", " "))
    cek("gagal kirim tidak menghapus data HP", "tetap di HP" in mob or "data tetap aman di HP" in mob)
    cek("ada penanda belum terkirim", "_sinkron" in mob and "belum terkirim" in mob)
    cek("sinkron ulang saat sinyal kembali", "addEventListener('online'" in mob)
    cek("sinkron ulang saat app dibuka", "await sinkronPenuh()" in mob)
    cek("pakai localStorage sebagai penyimpanan HP", "localStorage" in mob)
    cek("bisa ditarik dari HP tanpa server (mode file:)", "sedangMandiri" in mob
        and "location.protocol==='file:'" in mob.replace(" ", ""))
    cek("alamat server bisa diisi manual", "simpanServer" in mob and "K_SERVER" in mob)
    cek("antrean vs server tidak digandakan", "kunciServer" in mob)
    cek("semua kolom input ada (tanggal..keterangan)",
        all(k in mob for k in ('id="iTanggal"', 'id="iDriver"', 'id="iAlat"', 'id="iAktivitas"',
                               'id="iLokasi"', 'id="iLiter"', 'id="iHarga"', 'id="iKet"')))
    cek("sisa stok tidak diminta manual di HP", "SISA STOK" not in mob)
    cek("laporan WA berupa teks ringkasan", "teksLaporan" in mob and "terkirim" in mob.lower())
    # konsistensi id JS <-> HTML di halaman HP
    ids_js = set(re.findall(r"\$\('([A-Za-z0-9_]+)'\)", mob))
    ids_html = set(re.findall(r'id="([A-Za-z0-9_]+)"', mob))
    cek("semua id yang dipakai JS ada di HTML HP", not (ids_js - ids_html),
        f"hilang: {sorted(ids_js - ids_html)}")
    fungsi = set(re.findall(r"function\s+([A-Za-z0-9_]+)\s*\(", mob))
    dipanggil = set(re.findall(r"onclick=\"([A-Za-z0-9_]+)\(", mob))
    cek("semua onclick punya fungsinya", not (dipanggil - fungsi),
        f"tak terdefinisi: {sorted(dipanggil - fungsi)}")
    cek("sw.js rajin (cache baru + fallback offline)", "bbm-v3" in pathlib.Path(
        BASE / "static" / "sw.js").read_text(encoding="utf-8"))
    cek("manifest start_url /m", '"start_url": "/m"' in pathlib.Path(
        BASE / "static" / "manifest.webmanifest").read_text(encoding="utf-8"))

    # ---------- 11. rekap: catatan khusus & data terpotong
    print("\n11) Rekap: catatan khusus & data terpotong")
    # Di file Excel perusahaan, baris catatan punya tanggal asli tapi kolom
    # nama alat berisi teks "BBM HILANG/BOCOR" dan keterangannya ditandai
    # "catatan khusus". Bentuk itulah yang diuji di sini.
    idc = db.entry_simpan({"tanggal": "2026-09-19", "no_lambung": "BBM HILANG/BOCOR",
                           "lokasi": "PT GAS", "keluar": 3290, "harga": 12000,
                           "keterangan": "catatan khusus dari Excel: BBM HILANG/BOCOR"}, user_id=1)
    rk3 = db.rekap({})
    baris_c = [c for c in rk3["catatan"] if c["id"] == idc]
    cek("baris catatan terdeteksi", len(baris_c) == 1
        and abs(baris_c[0]["keluar"] - 3290) < 1,
        f'{len(baris_c)} baris, {[c["keluar"] for c in rk3["catatan"]]}')
    cek("liter catatan ikut dirinci", abs(rk3["catatan_keluar"]
        - sum(c["keluar"] for c in rk3["catatan"])) < 1, str(rk3["catatan_keluar"]))
    cek("pemakaian = total keluar dikurangi catatan",
        abs(rk3["pemakaian_keluar"] - (rk3["total_keluar"] - rk3["catatan_keluar"])) < 1,
        f'{rk3["pemakaian_keluar"]} vs {rk3["total_keluar"]}-{rk3["catatan_keluar"]}')
    cek("catatan tetap dihitung sisa stok",
        abs(rk3["sisa_stok"] - (rk3["stok_awal"] + rk3["total_masuk"] - rk3["total_keluar"])) < 1,
        str(rk3["sisa_stok"]))
    cek("catatan tidak ikut per_alat", all("HILANG" not in str(x["no_lambung"]).upper()
                                           for x in rk3["per_alat"]))
    cek("catatan tidak ikut per_driver", all(str(x["driver"] or "").strip()
                                             for x in rk3["per_driver"]))
    cek("catatan punya nama alat asli", baris_c[0]["no_lambung"] == "BBM HILANG/BOCOR")
    cek("catatan punya tanggal asli", baris_c[0]["tanggal"] == "2026-09-19",
        str(baris_c[0]["tanggal"]))
    # batas 30 baris: yang dipotong harus tetap dilaporkan, bukan hilang diam-diam
    for i in range(35):
        db.entry_simpan({"tanggal": "2026-09-20", "no_lambung": f"ALAT UJI {i:02d}",
                         "lokasi": "PT GAS", "keluar": 10, "harga": 17200}, user_id=1)
    rk4 = db.rekap({})
    cek("grup tampil dibatasi 30", len(rk4["per_alat"]) == 30, str(len(rk4["per_alat"])))
    sisa = rk4["sisa_alat"]
    cek("sisa grup dilaporkan", sisa["n"] > 0 and sisa["total_grup"] > 30,
        f'{sisa["n"]} dari {sisa["total_grup"]}')
    cek("jumlah liter 30 teratas + sisa = pemakaian",
        abs(sum(x["keluar"] for x in rk4["per_alat"]) + sisa["keluar"]
            - rk4["pemakaian_keluar"]) < 1, str(sisa["keluar"]))
    cek("per_driver punya biaya", all("biaya" in x for x in rk4["per_driver"]))
    cek("per_aktivitas ada", isinstance(rk4["per_aktivitas"], list))
    cek("harian punya tanggal/keluar/masuk",
        all({"tanggal", "keluar", "masuk"} <= set(h) for h in rk4["harian"]))
    # laporan & Excel harus ikut menyebut pemisahan catatan, bukan cuma rekap layar
    cek("laporan teks sebut catatan khusus", "Catatan khusus" in report.laporan_teks(rk3, db.get_settings(), {})
        and "Pemakaian alat" in report.laporan_teks(rk3, db.get_settings(), {}))
    cek("laporan html ada tabel catatan", "Catatan Khusus (bukan pemakaian alat)"
        in report.laporan_html(rk3, [], db.get_settings(), {}))
    cek("laporan tanpa catatan tidak bikin tabel kosong", "Catatan Khusus (bukan pemakaian alat)"
        not in report.laporan_html(db.rekap({"bulan": "2026-08"}), [], db.get_settings(), {}))
    wb_c = openpyxl.load_workbook(io.BytesIO(
        excel.ekspor_excel(db.semua_entries({}), rk3, db.get_settings()).getvalue()))
    cek("Excel punya sheet catatan", "CATATAN" in [s.upper() for s in wb_c.sheetnames],
        str(wb_c.sheetnames))
    # bersihkan supaya bagian lain tidak terpengaruh
    for r0 in db.entry_daftar({"alat": "ALAT UJI"})["rows"]:
        db.entry_hapus(r0["id"])
    for r0 in db.entry_daftar({"alat": "BBM HILANG"})["rows"]:
        db.entry_hapus(r0["id"])

    # --- sisa stok harus SALDO BERJALAN, bukan hitungan periode.
    # September (masuk 0, keluar 6.624) pernah tampil "sisa stok -6.624 L".
    cek("sisa stok bulan tanpa masuk tidak minus",
        db.rekap({"bulan": date.today().strftime("%Y-%m")})["sisa_stok"] > -1,
        str(db.rekap({"bulan": date.today().strftime("%Y-%m")})["sisa_stok"]))
    sem = db.rekap({})
    cek("sisa stok filter bulan = saldo akhir periode",
        abs(db.rekap({"bulan": "2026-09"})["sisa_stok"]
            - db.rekap({"sampai": "2026-09-30"})["sisa_stok"]) < 1,
        f'{db.rekap({"bulan": "2026-09"})["sisa_stok"]} vs '
        f'{db.rekap({"sampai": "2026-09-30"})["sisa_stok"]}')
    cek("sisa stok semua periode = total akhir",
        abs(sem["sisa_stok"] - (sem["stok_awal"] + sem["kumulatif_masuk"]
                                - sem["kumulatif_keluar"])) < 1, str(sem["sisa_stok"]))
    cek("sisa periode ini tetap dilaporkan terpisah",
        "sisa_periode_ini" in sem and abs(sem["sisa_periode_ini"]
                                          - (sem["stok_awal"] + sem["total_masuk"]
                                             - sem["total_keluar"])) < 1)

    print("\n" + "=" * 46)
    print(f"LULUS {LULUS} / GAGAL {GAGAL}")
    print("=" * 46)
    return 1 if GAGAL else 0


if __name__ == "__main__":
    sys.exit(main())
