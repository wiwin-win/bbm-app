"""Cek cepat server live: halaman, export, info LAN."""
import io
import os
import zipfile

import openpyxl
import requests

BASE = "http://127.0.0.1:8791"
s = requests.Session()
s.post(BASE + "/api/login", data={"username": "admin", "password": "admin"})

for p in ("/", "/m", "/static/manifest.webmanifest", "/static/sw.js", "/static/icon-192.png"):
    r = s.get(BASE + p)
    print(f"{p:34s} {r.status_code} {len(r.content)} bytes")

r = s.get(BASE + "/api/export/excel")
print("export excel:", r.status_code, len(r.content), "bytes | zip valid:", zipfile.is_zipfile(io.BytesIO(r.content)))
f = os.path.join(os.environ["LOCALAPPDATA"], "Temp", "bbm_export.xlsx")
with open(f, "wb") as fh:
    fh.write(r.content)
wb = openpyxl.load_workbook(f)
print("sheet:", wb.sheetnames, "| baris BBM:", wb["BBM"].max_row)

r = s.get(BASE + "/api/template/excel")
print("template:", r.status_code, len(r.content), "bytes")

info = s.get(BASE + "/api/ringkasan").json()
print("info LAN:", info.get("lan"), "| bulan:", info["data"]["bulan"])
