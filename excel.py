"""Import / export Excel untuk aplikasi Monitoring BBM & Alat."""
import io
import re
from datetime import date, datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import storage as db

KOLOM = [
    ("tanggal", "TANGGAL"),
    ("driver", "NAMA OPT / DRIVER"),
    ("no_lambung", "ALAT / NO LAMBUNG"),
    ("aktivitas", "DESKRIPSI AKTIVITAS"),
    ("lokasi", "LOKASI KEGIATAN"),
    ("keluar", "BBM / LTR (KELUAR)"),
    ("masuk", "BBM MASUK (LTR)"),
    ("harga", "HARGA / LTR (RP)"),
    ("hm_awal", "HM AWAL"),
    ("hm_akhir", "HM AKHIR"),
    ("keterangan", "KETERANGAN"),
]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
JUDUL_FONT = Font(bold=True, size=13, color="1F3864")
TOTAL_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
RP = '#,##0'
LTR = '#,##0.##'


def _rapi(header: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(header or "").upper())


def _cari_header(ws, maks: int = 15):
    """Cari baris header: baris yang memuat kata TANGGAL dan (LTR atau ALAT)."""
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=maks, values_only=True), start=1):
        teks = [_rapi(c) for c in row if c is not None]
        gab = "|".join(teks)
        if "TANGGAL" in gab and ("LTR" in gab or "ALAT" in gab or "LAMBUNG" in gab):
            return i, [str(c).strip() if c is not None else "" for c in row]
    return None, []


FMT_TGL = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%d.%m.%Y", "%d %b %Y", "%d %B %Y")


def _tgl_ketat(v):
    """Ubah nilai tanggal jadi YYYY-MM-DD, atau None kalau bukan tanggal.

    Dipakai saat impor: baris seperti 'BBM HILANG/BOCOR' di kolom tanggal
    harus ditolak, bukan dianggap tanggal.
    """
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date) and not isinstance(v, datetime):
        return v.isoformat()
    if isinstance(v, (int, float)):
        if 20000 <= float(v) <= 80000:          # serial tanggal Excel
            return db.tgl_iso(v)
        return None
    s = str(v).strip()
    for fmt in FMT_TGL:
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _peta_kolom(headers: list[str]) -> dict:
    """Peta field -> indeks kolom, cocokkan nama kolom secara longgar."""
    peta, dipakai = {}, set()
    for field, label in KOLOM:
        target = _rapi(label)
        for i, h in enumerate(headers):
            if i in dipakai or not h:
                continue
            hs = _rapi(h)
            if hs == target or target in hs or hs in target:
                peta[field] = i
                dipakai.add(i)
                break
    # cadangan manual untuk urutan kolom standar
    default = {"tanggal": 1, "driver": 2, "no_lambung": 3, "aktivitas": 4,
               "lokasi": 5, "keluar": 6, "masuk": 7}
    for k, v in default.items():
        if k not in peta and len(headers) > v:
            peta[k] = v
    return peta


def impor_excel(path: str, kosongkan: bool = False) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    # ---- cari sheet transaksi
    sheet, baris_head, headers, peta = None, 0, [], {}
    for nama in wb.sheetnames:
        h, hd = _cari_header(wb[nama])
        if h:
            sheet, baris_head, headers = nama, h, hd
            peta = _peta_kolom(headers)
            break
    if sheet is None:
        raise ValueError("Tidak menemukan baris header (TANGGAL / ALAT NO LAMBUNG) di file ini")
    ws = wb[sheet]

    baris_baru, driver_baru, alat_baru = [], set(), set()
    catatan_khusus: list[dict] = []
    tgl_terakhir = None
    for row in ws.iter_rows(min_row=baris_head + 1, values_only=True):
        if not any(c not in (None, "") for c in row):
            continue
        ambil = lambda f: row[peta[f]] if f in peta and len(row) > peta[f] else None  # noqa: E731
        nama_alat = str(ambil("no_lambung") or "").strip()
        if _rapi(nama_alat) in ("BBMHILANGBOCOR", "TOTAL", "SISASTOKAKHIRLTR"):
            continue
        tgl = _tgl_ketat(ambil("tanggal"))
        if tgl:
            tgl_terakhir = tgl
        if tgl is None:
            # Baris catatan: kolom tanggal berisi teks (mis. 'BBM HILANG/BOCOR').
            # Masih dicatat kalau ada nilai liter, pakai tanggal terakhir + keterangan.
            teks_tgl = str(ambil("tanggal") or "").strip()
            ada_liter = db._num(ambil("keluar")) or db._num(ambil("masuk"))
            if teks_tgl and ada_liter:
                rec_khusus = {
                    "tanggal": tgl_terakhir or date.today().isoformat(),
                    "driver": str(ambil("driver") or "").strip(),
                    "no_lambung": nama_alat or teks_tgl,
                    "aktivitas": str(ambil("aktivitas") or "").strip(),
                    "lokasi": str(ambil("lokasi") or "").strip(),
                    "keluar": db._num(ambil("keluar")),
                    "masuk": db._num(ambil("masuk")),
                    "harga": db._num(ambil("harga")) or db._num(db.get_settings().get("harga_default", 0)),
                    "keterangan": f"catatan khusus dari Excel: {teks_tgl}",
                    "jenis": "BBM",
                }
                baris_baru.append(rec_khusus)
                if nama_alat:
                    alat_baru.add(nama_alat)
                catatan_khusus.append({"tanggal_excel": teks_tgl, "no_lambung": nama_alat,
                                       "keluar": rec_khusus["keluar"], "masuk": rec_khusus["masuk"]})
            continue
        keterangan = str(ambil("keterangan") or "").strip()
        if not nama_alat:
            # Kolom alat kosong (biasanya baris BBM masuk / isi tangki).
            # Jangan pakai teks kolom tanggal sebagai nama alat — itu bikin
            # nama alat jadi tanggal. Beri nama yang jelas supaya laporan rapi.
            teks_tgl = str(ambil("tanggal") or "").strip()
            nama_alat = "BBM MASUK (TANPA ALAT)" if db._num(ambil("masuk")) else "TANPA ALAT"
            keterangan = (keterangan + f" | asal kolom tanggal: {teks_tgl}").strip(" |")
        rec = {
            "tanggal": tgl,
            "driver": str(ambil("driver") or "").strip(),
            "no_lambung": nama_alat,
            "aktivitas": str(ambil("aktivitas") or "").strip(),
            "lokasi": str(ambil("lokasi") or "").strip(),
            "keluar": db._num(ambil("keluar")),
            "masuk": db._num(ambil("masuk")),
            "harga": db._num(ambil("harga")) or db._num(db.get_settings().get("harga_default", 0)),
            "hm_awal": ambil("hm_awal") if "hm_awal" in peta else None,
            "hm_akhir": ambil("hm_akhir") if "hm_akhir" in peta else None,
            "keterangan": keterangan,
            "jenis": "BBM",
        }
        if not rec["keluar"] and not rec["masuk"]:
            continue
        if rec["driver"]:
            driver_baru.add(rec["driver"])
        if rec["no_lambung"]:
            alat_baru.add(rec["no_lambung"])
        baris_baru.append(rec)

    if not baris_baru:
        raise ValueError("Tidak ada baris data yang bisa diimpor dari sheet " + sheet)

    if kosongkan:
        with db.connect() as con:
            con.execute("DELETE FROM entries")
    for r in baris_baru:
        db.entry_simpan(r, user_id=1)

    n_drv = db.import_master("driver", [{"nama": d} for d in sorted(driver_baru)])
    n_alt = db.import_master("unit", [{"no_lambung": a, "vendor": "", "status": "Aktif"}
                                      for a in sorted(alat_baru)])
    # lokasi & aktivitas juga dikenali otomatis
    n_lok = db.import_master("location", [{"nama": x} for x in sorted(
        {r["lokasi"] for r in baris_baru if r["lokasi"]})])
    n_akt = db.import_master("activity", [{"nama": x} for x in sorted(
        {r["aktivitas"] for r in baris_baru if r["aktivitas"]})])

    # harga default ikut terisi dari data terakhir bila masih 0
    harga_terakhir = next((r["harga"] for r in reversed(baris_baru) if r["harga"]), 0)
    s = db.get_settings()
    if harga_terakhir and db._num(s.get("harga_default", 0)) == 0:
        db.set_settings({"harga_default": str(int(harga_terakhir))})
    t0, t1 = baris_baru[0]["tanggal"], baris_baru[-1]["tanggal"]
    return {"sheet": sheet, "baris": len(baris_baru), "dari": t0, "sampai": t1,
            "driver_baru": n_drv, "alat_baru": n_alt, "lokasi_baru": n_lok,
            "aktivitas_baru": n_akt, "kosongkan": kosongkan,
            "catatan_khusus": catatan_khusus}


def template_kosong() -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BBM"
    for i, (_, label) in enumerate(KOLOM, start=1):
        c = ws.cell(row=1, column=i, value=label)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(label) + 3)
    contoh = [datetime.today().strftime("%Y-%m-%d"), "CONTOH DRIVER", "LOADER (GAS 009)",
              "MAINTENANCE JALAN", "PT GAS", 150, 0, 17200, 17455, 17460, "contoh baris - hapus"]
    for i, v in enumerate(contoh, start=1):
        ws.cell(row=2, column=i, value=v)
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def _tgl_aman(tgl: str) -> tuple[str, str]:
    """Kembalikan (teks_tanggal, nama_bulan) yang aman dipakai di Excel."""
    s = str(tgl or "").strip()
    try:
        d = datetime.strptime(s, "%Y-%m-%d")
        return s, d.strftime("%B %Y")
    except ValueError:
        return s, ""


def ekspor_excel(rows: list[dict], rk: dict, setting: dict) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BBM"
    judul = setting.get("nama_instansi", "MONITORING BBM")
    ws["A1"] = f"PENCATATAN BAHAN BAKAR MINYAK (BBM) / SOLAR — {judul}"
    ws["A1"].font = JUDUL_FONT
    per = ""
    if rk.get("periode_awal"):
        per = f"Periode {rk['periode_awal']} s/d {rk['periode_akhir']}"
    ws["A2"] = per
    ws["A2"].font = Font(italic=True, size=10)

    # Kolom ditulis dari SATU daftar ini (label + cara ambil nilainya) supaya header
    # dan isi baris tidak mungkin bergeser posisi.
    def _nilai(kunci: str, r: dict, n: int, bulan: str):
        if kunci == "no":
            return n
        if kunci == "bulan":
            return bulan
        if kunci == "nilai_keluar":
            return r["keluar"] * r["harga"]
        if kunci == "nilai_masuk":
            return r["masuk"] * r["harga"]
        return r.get(kunci, "") or ""

    BARIS = [
        ("no", "NO", None),
        ("bulan", "BULAN", None),
        ("tanggal", "TANGGAL", None),
        ("driver", "NAMA OPT / DRIVER", None),
        ("no_lambung", "ALAT / NO LAMBUNG", None),
        ("aktivitas", "DESKRIPSI AKTIVITAS", None),
        ("lokasi", "LOKASI KEGIATAN", None),
        ("keluar", "BBM / LTR (KELUAR)", LTR),
        ("masuk", "BBM MASUK (LTR)", LTR),
        ("harga", "HARGA / LTR (RP)", RP),
        ("sisa_stok", "SISA STOK (LTR)", LTR),
        ("nilai_keluar", "NILAI BBM KELUAR (RP)", RP),
        ("nilai_masuk", "NILAI BBM MASUK (RP)", RP),
        ("keterangan", "KETERANGAN", None),
    ]
    hr = 3
    for i, (_, label, _fmt) in enumerate(BARIS, start=1):
        c = ws.cell(row=hr, column=i, value=label)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(26, len(label) + 3))
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)

    for n, r in enumerate(rows, start=1):
        _tgl, bulan = _tgl_aman(r["tanggal"])
        for i, (kunci, _label, fmt) in enumerate(BARIS, start=1):
            c = ws.cell(row=hr + n, column=i, value=_nilai(kunci, r, n, bulan))
            c.border = BORDER
            if fmt:
                c.number_format = fmt

    n = len(rows)
    tr = hr + n + 1
    kol = {label: i for i, (_, label, _f) in enumerate(BARIS, start=1)}
    ws.cell(row=tr, column=1, value="TOTAL")
    ws.cell(row=tr, column=kol["NO"], value=f"{n} baris")
    ws.cell(row=tr, column=kol["BBM / LTR (KELUAR)"], value=rk["total_keluar"]).number_format = LTR
    ws.cell(row=tr, column=kol["BBM MASUK (LTR)"], value=rk["total_masuk"]).number_format = LTR
    ws.cell(row=tr, column=kol["NILAI BBM KELUAR (RP)"], value=rk["total_biaya"]).number_format = RP
    for col in range(1, len(BARIS) + 1):
        ws.cell(row=tr, column=col).fill = TOTAL_FILL
        ws.cell(row=tr, column=col).font = Font(bold=True)

    # ---- sheet REKAP
    wr = wb.create_sheet("REKAP")
    wr.column_dimensions["A"].width = 34
    for col in "BCDE":
        wr.column_dimensions[col].width = 18
    wr["A1"] = "REKAP OTOMATIS"
    wr["A1"].font = JUDUL_FONT
    isi = [("Stok Awal (Ltr)", rk["stok_awal"]),
           ("Total BBM Masuk (Ltr)", rk["total_masuk"]),
           ("Total BBM Keluar (Ltr)", rk["total_keluar"]),
           ("Sisa Stok Akhir (Ltr)", rk["sisa_stok"]),
           ("Total Biaya BBM Keluar (Rp)", rk["total_biaya"]),
           ("Harga Rata-rata (Rp/Ltr)", rk["harga_rata"]),
           ("Jumlah Baris Data", rk["baris"])]
    # Baris catatan (mis. "BBM HILANG/BOCOR") mengurangi stok tapi bukan pemakaian
    # alat. Tanpa baris ini, angka total_keluar dan jumlah rincian tidak nyambung.
    if rk.get("catatan"):
        isi += [("Pemakaian Alat Saja (Ltr)", rk.get("pemakaian_keluar", 0)),
                ("Biaya Pemakaian Alat (Rp)", rk.get("pemakaian_biaya", 0)),
                ("Catatan Khusus (Ltr)", rk.get("catatan_keluar", 0)),
                ("Biaya Catatan Khusus (Rp)", rk.get("catatan_biaya", 0)),
                ("Jumlah Baris Catatan", len(rk["catatan"]))]
    for i, (k, v) in enumerate(isi, start=3):
        wr.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = wr.cell(row=i, column=2, value=v)
        c.number_format = RP if "Rp" in k or "Harga" in k else LTR

    def tbl(judul, data, kolom, kolom_label, mulai, sisa=None):
        if sisa and sisa.get("n"):
            judul += f" — {sisa.get('n')} entri lain ({sisa.get('keluar', 0)} L) tidak ditampilkan"
        wr.cell(row=mulai, column=1, value=judul).font = Font(bold=True, color="1F3864")
        for j, lab in enumerate(kolom_label, start=1):
            c = wr.cell(row=mulai + 1, column=j, value=lab)
            c.fill = HEADER_FILL
            c.font = HEADER_FONT
        for i, r in enumerate(data, start=1):
            for j, k in enumerate(kolom, start=1):
                c = wr.cell(row=mulai + 1 + i, column=j, value=r.get(k, ""))
                if k in ("keluar", "biaya"):
                    c.number_format = LTR if k == "keluar" else RP
        # baris total: angka total grup asli, bukan cuma yang tampil
        tot = mulai + len(data) + 1
        wr.cell(row=tot, column=1, value="TOTAL").font = Font(bold=True)
        for j, k in enumerate(kolom, start=1):
            if k in ("keluar", "biaya"):
                nilai = sum(r.get(k, 0) or 0 for r in data) + (
                    (sisa or {}).get(k, 0) if sisa else 0)
                c = wr.cell(row=tot, column=j, value=nilai)
                c.number_format = LTR if k == "keluar" else RP
                c.font = Font(bold=True)
        return tot + 2

    pos = 3 + len(isi) + 2
    pos = tbl("BBM KELUAR PER ALAT / NO LAMBUNG", rk["per_alat"], ["no_lambung", "keluar", "biaya"],
              ["ALAT / NO LAMBUNG", "LITER", "BIAYA (RP)"], pos, rk.get("sisa_alat"))
    pos = tbl("BBM PER LOKASI KEGIATAN", rk["per_lokasi"], ["lokasi", "keluar", "biaya"],
              ["LOKASI", "LITER", "BIAYA (RP)"], pos, rk.get("sisa_lokasi"))
    pos = tbl("BBM PER DRIVER (30 TERATAS)", rk["per_driver"], ["driver", "keluar", "biaya"],
              ["NAMA OPT / DRIVER", "LITER", "BIAYA (RP)"], pos, rk.get("sisa_driver"))
    pos = tbl("BBM PER AKTIVITAS", rk["per_aktivitas"], ["aktivitas", "keluar", "biaya"],
              ["DESKRIPSI AKTIVITAS", "LITER", "BIAYA (RP)"], pos, rk.get("sisa_aktivitas"))

    # ---- sheet CATATAN (kalau ada): baris yang bukan pemakaian alat
    if rk.get("catatan"):
        wc = wb.create_sheet("CATATAN")
        wc["A1"] = "CATATAN KHUSUS — BUKAN PEMAKAIAN ALAT"
        wc["A1"].font = JUDUL_FONT
        wc["A2"] = ("Baris ini tetap mengurangi sisa stok, tapi tidak dihitung sebagai "
                    "pemakaian alat/driver supaya rekap tidak menyesatkan.")
        wc["A2"].font = Font(italic=True, size=10)
        for i, lab in enumerate(["KETERANGAN ASLI", "TANGGAL DATA", "LOKASI",
                                 "LITER (KELUAR)", "BIAYA (RP)"], start=1):
            c = wc.cell(row=4, column=i, value=lab)
            c.fill = HEADER_FILL
            c.font = HEADER_FONT
            wc.column_dimensions[get_column_letter(i)].width = 22 if i == 1 else 16
        for n, c0 in enumerate(rk["catatan"], start=5):
            wc.cell(row=n, column=1, value=c0.get("no_lambung") or "")
            wc.cell(row=n, column=2, value=c0.get("tanggal") or "")
            wc.cell(row=n, column=3, value=c0.get("lokasi") or "")
            wc.cell(row=n, column=4, value=c0.get("keluar") or 0).number_format = LTR
            wc.cell(row=n, column=5, value=c0.get("biaya") or 0).number_format = RP
        tot = 5 + len(rk["catatan"])
        wc.cell(row=tot, column=3, value="TOTAL").font = Font(bold=True)
        wc.cell(row=tot, column=4, value=rk.get("catatan_keluar", 0)).number_format = LTR
        wc.cell(row=tot, column=5, value=rk.get("catatan_biaya", 0)).number_format = RP
        wc.cell(row=tot + 2, column=1, value="Pemakaian alat saja (Ltr)").font = Font(bold=True)
        wc.cell(row=tot + 2, column=2, value=rk.get("pemakaian_keluar", 0)).number_format = LTR

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio
