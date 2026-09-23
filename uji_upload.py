"""Uji alur "perbarui dari Excel" dan "unduh ke Excel" pada server hidup.

Skrip ini: bikin file Excel uji -> unggah lewat API -> cek datanya masuk ->
unduh Excel hasil ekspor -> buka lagi dengan openpyxl -> pastikan isinya benar
-> hapus data uji supaya database bersih.

    python uji_upload.py          (server harus jalan di 127.0.0.1:8791)
"""
import io
import sys
from datetime import date

import requests
from openpyxl import Workbook, load_workbook

BASE = "http://127.0.0.1:8791"
PENANDA = "UJI UNGGAH"
GAGAL = 0


def cek(nama, syarat, info=""):
    global GAGAL
    if not syarat:
        GAGAL += 1
    print(f"  [{'OK  ' if syarat else 'GAGAL'}] {nama} {info}")


def bikin_excel() -> bytes:
    """Excel uji dengan header seperti file asli perusahaan."""
    wb = Workbook()
    ws = wb.active
    ws.title = "BBM"
    ws.append(["TANGGAL", "NAMA OPT / DRIVER", "ALAT NO LAMBUNG", "DESKRIPSI AKTIVITAS",
               "LOKASI", "BBM/LTR KELUAR", "BBM MASUK", "HARGA/LTR", "KETERANGAN"])
    ws.append(["2026-09-22", PENANDA, "UJI-ALAT-01", "UJI AKTIVITAS", "UJI LOKASI", 12.5, "", 17200, "uji"])
    ws.append(["2026-09-22", PENANDA, "UJI-ALAT-02", "UJI AKTIVITAS", "UJI LOKASI", 7.5, "", 17200, ""])
    ws.append(["2026-09-22", "", "", "UJI MASUK", "UJI LOKASI", "", 1000, 16000, "isi tangki"])
    ws.append(["2026-09-22", "UJI HILANG", "", "UJI BARIS TANPA ALAT", "UJI LOKASI", 5, "", 17200, ""])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


print("=" * 52)
print("UJI UNGGAH / UNDUH EXCEL (server hidup)")
print("=" * 52)

s = requests.Session()
r = s.post(BASE + "/api/login", data={"username": "admin", "password": "admin"})
cek("login admin", r.status_code == 200, f"HTTP {r.status_code}")
if r.status_code != 200:
    sys.exit(1)

sebelum = s.get(BASE + "/api/entries?limit=1").json()["total"]
print(f"  (data di database sebelum uji: {sebelum} baris)")

# ---------- 1. unggah
print("\n1) Perbarui data dari Excel (unggah)")
isian = bikin_excel()
r = s.post(BASE + "/api/import/excel",
           files={"file": ("uji.xlsx", isian,
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
           data={"kosongkan": "0"})
cek("unggah diterima", r.status_code == 200, f"HTTP {r.status_code} {r.text[:120]}")
hasil = r.json()
print(f"       hasil: {hasil}")
cek("terbaca 4 baris", hasil.get("baris") == 4, f"baris={hasil.get('baris')}")
cek("tidak mengosongkan data lama", hasil.get("kosongkan") is False)
cek("periode terbaca", hasil.get("dari") == "2026-09-22", f"{hasil.get('dari')} s/d {hasil.get('sampai')}")

sesudah = s.get(BASE + "/api/entries?limit=1").json()["total"]
cek("jumlah data bertambah 4", sesudah == sebelum + 4, f"{sebelum} -> {sesudah}")

baru = s.get(BASE + "/api/entries",
             params={"dari": "2026-09-22", "sampai": "2026-09-22", "limit": 50}).json()["rows"]
cek("4 baris uji ada di database", len(baru) == 4, f"{len(baru)} baris")
semua = baru
alat = {x["no_lambung"] for x in semua if x["no_lambung"]}
cek("no lambung terbaca", "UJI-ALAT-01" in alat and "UJI-ALAT-02" in alat, str(alat))
cek("baris tanpa alat tidak dinamai tanggal",
    all(not str(a)[:4].isdigit() for a in alat), str(sorted(alat)))
keluar = sum(float(x["keluar"] or 0) for x in semua)
masuk = sum(float(x["masuk"] or 0) for x in semua)
cek("liter keluar terbaca benar (25)", abs(keluar - 25) < 0.01, str(keluar))
cek("liter masuk terbaca benar (1000)", abs(masuk - 1000) < 0.01, str(masuk))
master = s.get(BASE + "/api/master").json()
cek("master alat ikut bertambah", any(u["no_lambung"] == "UJI-ALAT-01" for u in master["unit"]))

# ---------- 2. unduh
print("\n2) Unduh data ke Excel")
r = s.get(BASE + "/api/export/excel", params={"dari": "2026-09-22", "sampai": "2026-09-22"})
cek("unduh berhasil", r.status_code == 200 and len(r.content) > 4000,
    f"HTTP {r.status_code} {len(r.content)} bytes")
cek("bertipe xlsx", "spreadsheetml" in r.headers.get("Content-Type", ""))
cek("nama file ada tanggal", "BBM_" in r.headers.get("Content-Disposition", ""))

wb = load_workbook(io.BytesIO(r.content))
cek("file bisa dibuka Excel/openpyxl", wb.sheetnames is not None, str(wb.sheetnames))
cek("ada sheet BBM", "BBM" in wb.sheetnames, str(wb.sheetnames))
cek("ada sheet REKAP", "REKAP" in wb.sheetnames, str(wb.sheetnames))
isinya = "\n".join(str(c.value) for row in wb["BBM"].iter_rows() for c in row if c.value is not None)
cek("data uji ikut terunduh", PENANDA in isinya)
cek("liter 12.5 ikut terunduh", "12.5" in isinya)

# ---------- 3. template
print("\n3) Template Excel kosong")
r = s.get(BASE + "/api/template/excel")
wb2 = load_workbook(io.BytesIO(r.content))
kepala = [c.value for c in wb2[wb2.sheetnames[0]][1] if c.value]
cek("template berisi baris judul kolom", len(kepala) >= 5, str(kepala[:6]))

# ---------- 4. bersihkan
print("\n4) Bersihkan data uji")
for x in semua:
    s.delete(BASE + f"/api/entries/{x['id']}")
for u in master["unit"]:
    if str(u["no_lambung"]).startswith("UJI-ALAT") or str(u["no_lambung"]).startswith("TANPA ALAT") \
            or str(u["no_lambung"]).startswith("BBM MASUK (TANPA"):
        s.delete(BASE + f"/api/master/unit/{u['id']}")
for d in master["driver"]:
    if d["nama"] in (PENANDA, "UJI HILANG"):
        s.delete(BASE + f"/api/master/driver/{d['id']}")
for l in master["location"]:
    if str(l["nama"]).startswith("UJI "):
        s.delete(BASE + f"/api/master/location/{l['id']}")
for a in master["activity"]:
    if str(a["nama"]).startswith("UJI "):
        s.delete(BASE + f"/api/master/activity/{a['id']}")
akhir = s.get(BASE + "/api/entries?limit=1").json()["total"]
cek("database kembali seperti semula", akhir == sebelum, f"{akhir} baris (harus {sebelum})")
sisa = s.get(BASE + "/api/entries", params={"q": PENANDA, "limit": 5}).json()["rows"]
cek("tidak ada sisa baris uji", not sisa, f"{len(sisa)} sisa")

print("\n" + "=" * 52)
print("GAGAL" if GAGAL else "UNGGAH & UNDUH EXCEL BERJALAN", f"({GAGAL} gagal)")
print("=" * 52)
sys.exit(1 if GAGAL else 0)
